"""Rebuild and audit the 15-sector IC accounting extracts from supplied workbooks.

Input workbooks remain untouched.  All figures here are current-price workbook
amounts in 万元; the competitive-use table does not identify domestic users.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(os.environ.get("IC_ECONGYM_SOURCE_DIR", Path(__file__).resolve().parents[3]))
OUT = Path(os.environ.get("IC_ECONGYM_IO15_DIR", ROOT / "outputs" / "io_15_2018_2020_2023"))
FILES = {
    2018: ROOT / "终-2018投入产出表(1).xlsx",
    2020: ROOT / "终-2020投入产出表(1).xlsx",
    2023: ROOT / "Q-2023投入产出表(1).xlsx",
}
GROUPS = ["设计与关键投入"] * 7 + ["芯片制造与封测"] * 2 + ["系统集成与终端应用"] * 6
META_COLUMNS = ["中间使用", "消费支出", "固定资本形成", "存货变动", "出口", "最终使用", "进口", "总产出"]


def clean_label(value: object) -> str:
    return "".join(str(value).split())


def write_csv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def number(value: object, where: str) -> float:
    if not isinstance(value, (int, float)) or not np.isfinite(value):
        raise ValueError(f"Missing or nonnumeric cached value: {where}: {value!r}")
    return float(value)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    all_data = {}
    source_meta = {}
    sector_names = None
    metrics = []
    sectors = []
    flows = []
    audit = {}
    for year, path in FILES.items():
        book = openpyxl.load_workbook(path, data_only=True, read_only=False)
        sheet = book["完表"]
        if clean_label(sheet["A1"].value) != "单位：万元":
            raise ValueError(f"Unexpected unit in {path.name}")
        names = [clean_label(sheet.cell(4, c).value) if c <= 20 else clean_label(sheet.cell(3, c).value)
                 for c in range(6, 63)]
        row_names = [clean_label(sheet.cell(r, 4).value) if r <= 19 else clean_label(sheet.cell(r, 3).value)
                     for r in range(5, 62)]
        if len(names) != 57 or names != row_names or any(not name or name == "None" for name in names):
            raise ValueError(f"57-sector row/column labels do not align: {path.name}")
        if sector_names is None:
            sector_names = names
        elif names != sector_names:
            raise ValueError(f"Labels changed across years: {path.name}")
        z = np.array([[number(sheet.cell(r, c).value, f"{path.name}!{sheet.cell(r,c).coordinate}")
                       for c in range(6, 63)] for r in range(5, 62)], dtype=np.float64)
        use = np.array([[number(sheet.cell(r, c).value, f"{path.name}!{sheet.cell(r,c).coordinate}")
                         for c in range(63, 71)] for r in range(5, 62)], dtype=np.float64)
        va = np.array([number(sheet.cell(67, c).value, f"{path.name}!{sheet.cell(67,c).coordinate}")
                       for c in range(6, 63)], dtype=np.float64)
        x = use[:, 7]
        if np.any(x <= 0) or np.any(z < -1e-8):
            raise ValueError(f"Nonpositive output or negative intermediate use: {path.name}")
        row_residual = z.sum(axis=1) + use[:, 5] - use[:, 6] - x
        column_residual = z.sum(axis=0) + va - x
        use_residual = z.sum(axis=1) - use[:, 0]
        final_residual = use[:, 1:5].sum(axis=1) - use[:, 5]
        va_residual = np.array([sum(number(sheet.cell(r, c).value, f"{path.name}!{sheet.cell(r,c).coordinate}")
                                    for r in range(63, 67)) - va[c-6] for c in range(6, 63)])
        max_abs = {"row": float(np.max(np.abs(row_residual))),
                   "column": float(np.max(np.abs(column_residual))),
                   "intermediate": float(np.max(np.abs(use_residual))),
                   "final": float(np.max(np.abs(final_residual))),
                   "value_added": float(np.max(np.abs(va_residual)))}
        max_relative = max(max_abs.values()) / float(x.max())
        if max(max_abs.values()) > 1e-3 or max_relative > 1e-10:
            raise ValueError(f"Accounting identities fail: {year}: {max_abs}")
        intra = z[:15, :15]
        a15 = intra / x[None, :15]
        eig = float(max(abs(np.linalg.eigvals(a15))))
        if eig >= 1:
            raise ValueError(f"Intra-15 technical matrix is not productive: {year}: {eig}")
        inverse = np.linalg.inv(np.eye(15) - a15)
        multiplier = inverse.sum(axis=0)
        np.savez_compressed(OUT / f"io57_{year}.npz", z_competitive=z, final_use=use[:, 1:5],
                            imports=use[:, 6], gross_output=x, value_added=va,
                            ic15_intermediate=intra, ic15_coefficients=a15,
                            ic15_inverse=inverse)
        write_csv(OUT / f"ic15_intermediate_{year}.csv", ["供给部门\\使用部门", *sector_names[:15]],
                  [[sector_names[i], *[float(v) for v in intra[i]]] for i in range(15)])
        write_csv(OUT / f"ic15_coefficients_{year}.csv", ["供给部门\\使用部门", *sector_names[:15]],
                  [[sector_names[i], *[float(v) for v in a15[i]]] for i in range(15)])
        for i in range(15):
            sectors.append([year, i + 1, GROUPS[i], sector_names[i], x[i], va[i], va[i] / x[i],
                            use[i, 0], use[i, 1], use[i, 2], use[i, 3], use[i, 4], use[i, 5],
                            use[i, 6], use[i, 6] / x[i], intra[:, i].sum(), intra[i].sum(),
                            a15[:, i].sum(), multiplier[i]])
        for gi, group_from in enumerate(dict.fromkeys(GROUPS)):
            for gj, group_to in enumerate(dict.fromkeys(GROUPS)):
                a = [i for i, group in enumerate(GROUPS) if group == group_from]
                b = [j for j, group in enumerate(GROUPS) if group == group_to]
                flows.append([year, group_from, group_to, float(z[np.ix_(a, b)].sum())])
        metrics.append([year, len(names), 15, 42, float(x[:15].sum()), float(va[:15].sum()),
                        float(use[:15, 6].sum()), float(use[:15, 5].sum()), float(intra.sum()),
                        float(z[:15, 15:].sum()), float(z[15:, :15].sum()), eig,
                        float(multiplier.mean()), max_abs["row"], max_abs["column"]])
        audit[str(year)] = {"source": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                            "sheet": "完表", "sector_rows": "5:61", "sector_columns": "F:BJ",
                            "ic_rows": "5:19", "ic_columns": "F:T", "unit": "万元",
                            "max_absolute_residual_wanyuan": max_abs,
                            "max_relative_residual": max_relative,
                            "cached_values_used": True}
        all_data[year] = {"x": x, "va": va, "imports": use[:, 6], "final": use[:, 5]}
    write_csv(OUT / "sector_dictionary.csv", ["sector_id", "scope", "group", "sector_label"],
              [[i + 1, "IC15" if i < 15 else "external42", GROUPS[i] if i < 15 else "外部部门", sector_names[i]]
               for i in range(57)])
    write_csv(OUT / "metrics_by_year.csv",
              ["year", "all_sector_count", "ic_sector_count", "external_sector_count", "ic_gross_output_wanyuan",
               "ic_value_added_wanyuan", "ic_import_wanyuan", "ic_final_use_wanyuan",
               "ic_to_ic_intermediate_wanyuan", "ic_to_external_intermediate_wanyuan",
               "external_to_ic_intermediate_wanyuan", "ic15_coefficient_spectral_radius",
               "mean_intra15_leontief_multiplier", "row_balance_max_abs_wanyuan", "column_balance_max_abs_wanyuan"], metrics)
    write_csv(OUT / "sector_metrics.csv",
              ["year", "sector_id", "group", "sector_label", "gross_output_wanyuan", "value_added_wanyuan",
               "value_added_ratio", "intermediate_use_wanyuan", "consumption_wanyuan", "fixed_capital_wanyuan",
               "inventory_change_wanyuan", "exports_wanyuan", "final_use_wanyuan", "imports_wanyuan",
               "imports_to_output_ratio", "ic_inputs_wanyuan", "sales_to_ic_wanyuan",
               "intra15_direct_input_coefficient_sum", "intra15_leontief_column_sum"], sectors)
    write_csv(OUT / "group_flows.csv", ["year", "supplier_group", "user_group", "intermediate_wanyuan"], flows)
    changes = []
    for old, new in [(2018, 2020), (2020, 2023), (2018, 2023)]:
        for i in range(15):
            for metric, key in [("gross_output", "x"), ("value_added", "va"), ("imports", "imports"),
                                ("final_use", "final")]:
                before, after = float(all_data[old][key][i]), float(all_data[new][key][i])
                changes.append([old, new, i + 1, sector_names[i], metric, before, after,
                                after - before, (after / before - 1) if before else ""])
    write_csv(OUT / "sector_changes.csv", ["from_year", "to_year", "sector_id", "sector_label",
                                           "metric", "from_wanyuan", "to_wanyuan", "change_wanyuan",
                                           "nominal_change_ratio"], changes)
    (OUT / "audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"years": list(FILES), "sectors": 15, "max_residual":
                      max(max(v["max_absolute_residual_wanyuan"].values()) for v in audit.values())}, ensure_ascii=False))


if __name__ == "__main__":
    main()
